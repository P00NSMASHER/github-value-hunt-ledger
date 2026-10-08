"""Rendered PDF / DOCX consistency and source-identity checks."""
from pathlib import Path
import hashlib,json,re,sys
from PIL import Image,ImageOps,ImageDraw
import fitz
from docx import Document
P=Path(__file__).resolve().parent.parent
truth=json.loads((P/'source/commercial_truth.json').read_text())
reports=[];page_imgs=[];error=[]
for file in sorted((P/'pdf').glob('*.pdf')):
    word=P/'docx'/file.with_suffix('.docx').name
    if not word.exists():error.append('missing DOCX for '+file.name);continue
    pdf=fitz.open(file);pages=len(pdf)
    doc=Document(word)
    words='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    txt='\n'.join(p.get_text() for p in pdf)
    no_match=[];all_amounts=set(re.findall(r'\$[\d,]+\.\d\d',words))
    for x in sorted(all_amounts):
        if x not in txt:no_match.append(x)
    if no_match:error.append(file.name+' currency not in pdf '+str(no_match))
    if 'M2C v1.0' not in txt:error.append(file.name+' version not in PDF')
    for number,page in enumerate(pdf,1):
        if round(page.rect.width)!=612 or round(page.rect.height)!=792:error.append(file.name+' non US Letter')
        out=[];tiny=[]
        for line in page.get_text('dict')['blocks']:
            if 'lines' not in line:continue
            for item in line['lines']:
                for span in item['spans']:
                    if not span['text'].strip():continue
                    bb=fitz.Rect(span['bbox']);
                    if bb.x0<4 or bb.y0<3 or bb.x1>page.rect.width-4 or bb.y1>page.rect.height-3:out.append(span['text'][:35])
                    if span['size']<6.9:tiny.append((span['size'],span['text'][:30]))
        if out:error.append(file.name+':'+str(number)+' overflow '+str(out))
        if tiny:error.append(file.name+':'+str(number)+' tiny '+str(tiny))
        if number==1:
            gray=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),colorspace=fitz.csGRAY)
            Image.frombytes('L',[gray.width,gray.height],gray.samples).save(P/'qa'/('grayscale_'+file.stem+'.png'))
        pix=page.get_pixmap(matrix=fitz.Matrix(1.35,1.35),alpha=False)
        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        dest=P/'qa'/f'{file.stem}_page{number}.png'; im.save(dest)
        im.thumbnail((480,620));page_imgs.append((f'{file.stem} {number}/{pages}',im.copy()))
    reports.append(dict(file=file.name,pages=pages,currency_tokens=len(all_amounts),pdf_sha256=hashlib.sha256(file.read_bytes()).hexdigest(),docx_sha256=hashlib.sha256(word.read_bytes()).hexdigest(),parity_money=not no_match))
for i in range(0,len(page_imgs),6):
    group=page_imgs[i:i+6];sheet=Image.new('RGB',(3*515,2*690),'#e9eeeb');dr=ImageDraw.Draw(sheet)
    for j,(label,im) in enumerate(group):
        x=(j%3)*515+18;y=(j//3)*690+26;sheet.paste(im,(x,y));dr.text((x,y-18),label[:62],fill='#17332e')
    sheet.save(P/'qa'/f'm2c_sheet_{i//6+1}.jpg',quality=90)
for kind,name in [('wordmark','retally-wordmark-approved.png'),('emblem','retally-emblem-approved.png')]:
    actual=hashlib.sha256((P/'assets'/name).read_bytes()).hexdigest()
    if actual!=truth['approved_brand'][kind+'_sha256']:error.append('MASTER ART HASH '+kind)
result={'documents':reports,'total_pages':sum(x['pages'] for x in reports),'errors':error,'visual_review':'contact sheets generated; inspect before approval','approved_source_sha_pass':not any('MASTER ART' in s for s in error)}
(P/'qa'/'m2c_acceptance.json').write_text(json.dumps(result,indent=2)+'\n')
print('PDFS',len(reports),'PAGES',result['total_pages'],'ERRORS',len(error))
for e in error:print('ERROR',e)
for x in reports: print(x['file'],x['pages'],'money parity',x['parity_money'])
sys.exit(bool(error))