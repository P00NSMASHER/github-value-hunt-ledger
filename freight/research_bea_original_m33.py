"""RETALLY Lab 14 Mission 33: Original BEA GrossOutput.xlsx publication replay.

Unlike prior M32's verified DBnomics redistribution, this module
authenticates an ORIGINAL official BEA publisher XLSX from apps.bea.gov,
validating its full SHA-256 and byte length, table identity, date, industry,
NAICS, 8 annual quantities, price indices, nominal gross output and
chained-dollar consistency. It does NOT authenticate any real customer invoice
billing mistakes, refunds, buyer demand or recoverable market TAM.
"""
from __future__ import annotations

import argparse
from datetime import date
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from hashlib import sha1, sha256
from html import escape
from io import BytesIO
import json
from pathlib import Path
from urllib.request import Request, urlopen

from freight.research_evidence_gate import canon_hash
from freight.research_bea_gross_output_m32 import (
    report_shell as m32_redistributor_report,
    OFFLINE as M32_OFFLINE,
)
from freight.research_bls_sectoral_output_m31 import (
    report_shell as m31_sectoral_output_report,
    OFFLINE as M31_OFFLINE,
)

BOOK = Path(__file__).parent / "research/BEA_ORIGINAL_ACCESS_M33_20261012.json"
BOOK_SHA1 = "9a26ba7812b6569877c71adbd5cb8ddb430c2071"
BEA_URL = "https://apps.bea.gov/industry/Release/XLS/GDPxInd/GrossOutput.xlsx"
BEA_ORIGINAL_SHA256 = "a5a88e808bd88c5b48d7c1f572acb4d9b113d0d753828ed8b45d8800c1eb1f46"
BEA_ORIGINAL_BYTES = 506819
SCOPE="LAB14_M33_DIRECT_ORIGINAL_BEA_GROSS_OUTPUT_XLSX_TABLE16_NAICS484_SOURCE"
BOOK_SCOPE="LAB14_M33_DIRECT_FIRST_PARTY_BEA_GROSSOUTPUT_XLSX_484_TABLE16_EIGHT_YEAR_PROOF"
MIN_DATE=date(2026,10,12)
YEARS=tuple(range(2017,2025))
OFFLINE="M33_OFFLINE_PREVIOUS_BEA_DIRECT_XLSX_SOURCE_VALUES_NOT_REPLAYED_THIS_RUN"
VERIFIED="M33_ORIGINAL_BEA_506819_BYTE_GROSSOUTPUT_XLSX_4_TABLES_32_YEAR_CELLS_VERIFIED"
BLOCKED="M33_ORIGINAL_BEA_XLSX_NETWORK_UNAVAILABLE_NO_NEW_SOURCE_REPLAY"
SHEETS=("TGO103-A","TGO104-A","TGO105-A","TGO106-A")
EXPECTED_SHEETS=(
    "Contents","TGO101-A","TGO101-Q","TGO103-A","TGO103-Q",
    "TGO104-A","TGO104-Q","TGO105-A","TGO105-Q",
    "TGO106-A","TGO106-Q","TGO107-A","TGO107-Q",
)
TITLES={
    "TGO103-A":"Chain-Type Quantity Indexes for Gross Output by Industry",
    "TGO104-A":"Chain-Type Price Indexes for Gross Output by Industry",
    "TGO105-A":"Gross Output by Industry",
    "TGO106-A":"Real Gross Output by Industry",
}
UNITS={
    "TGO103-A":"[2017=100]",
    "TGO104-A":"[2017=100]",
    "TGO105-A":"[Millions of dollars]",
    "TGO106-A":"[Millions of 2017 chain dollars]",
}
EXPECTED={
    "TGO103-A":("100.000","102.952","102.398","103.678",
                 "109.538","108.519","107.307","108.047"),
    "TGO104-A":("100.000","106.173","108.901","108.653",
                 "121.921","144.891","138.837","137.494"),
    "TGO105-A":("367154","401324","409425","413595",
                 "490334","577292","546993","545435"),
    "TGO106-A":("367154","377992","375959","380657",
                 "402174","398432","393982","396699"),
}
MAX_BYTES=1_000_000
SOURCE_REVENUE_UNIT="MILLIONS_NOMINAL_USD_NOT_CUSTOMER_FREIGHT_INVOICES"
FALSE_GUARDS=(
    "independent_original_bea_source_proved_by_official_web_landing",
    "first_party_metadata_pages_prove_numeric_table_values",
    "http200_bea_api_error_means_auth_success",
    "unauthenticated_api_error_means_series_does_not_exist",
    "redistributor_values_are_first_party_bea_authenticated",
    "unauthorized_api_userid_or_registration_attempts_permitted",
    "be_a_gdp_value_added_can_substitute_for_gross_output",
    "original_bea_484_2017_2024_numeric_verification_completed",
    "gross_output_vs_sectoral_difference_is_validated_model_error",
    "gross_output_vs_ppi_proxy_difference_is_validated_model_error",
    "original_data_same_as_customer_invoices",
    "bea_index_is_direct_ltl_484122_recovery_market",
    "census_revenue_ppi_proxy_is_official_real_output",
    "bls_real_sectoral_output_equivalent_to_bea_gross_output",
    "direct_bea_file_does_not_certify_recoverable_freight_market",
)
# Forbid anyone from setting anything that implies real-money recovery.
SOURCE_NO_RECOVERY={
    "actual_real_customer_invoice_rows":0,
    "actual_independently_confirmed_invoice_overcharges":0,
    "real_refunds_usd":None,
    "recoverable_market_tam_usd":None,
    "actual_buyer_demand_established":False,
    "external_contact_or_marketing_claim_authorized":False,
    "private_customer_or_carrier_documents_accessed":False,
    "eighteen_independent_customer_pilot_approvals_satisfied":False,
    "existing_lab_count_changed":False,
    "scheduled_tasks_modified":False,
}


def _git_blob(raw:bytes)->str:
    return sha1(b"blob "+str(len(raw)).encode("ascii")+b"\0"+raw).hexdigest()


def _book()->dict:
    raw=BOOK.read_bytes()
    if _git_blob(raw)!=BOOK_SHA1:
        raise ValueError("M33_FROZEN_BEA_PUBLISHER_SOURCE_CATALOG_BLOB_DRIFT")
    return json.loads(raw.decode("utf-8"))


def _decimal(value:object)->Decimal:
    if type(value) not in (int,float,str,Decimal):
        raise ValueError("M33_ORIGINAL_BEA_CELL_NOT_NUMERIC")
    try:
        d=Decimal(str(value))
    except InvalidOperation as e:
        raise ValueError("M33_BEA_ORIGINAL_WORKBOOK_CELL_NONDECIMAL") from e
    if not d.is_finite() or d<=0 or d>=100000000:
        raise ValueError("M33_BEA_ORIGINAL_WORKBOOK_INDEX_OR_MILLIONS_OUT_OF_RANGE")
    return d


def validate_book(b:object,*,as_of:date)->list[str]:
    try:
        if type(b) is not dict or b!=_book():
            return ["M33_ORIGINAL_BEA_SOURCE_CATALOG_CHANGED_OR_UNAUTHORIZED_FIELDS"]
        errors=[]
        if type(as_of) is not date or as_of<MIN_DATE:
            errors.append("M33_BEFORE_OFFICIAL_BEA_SOURCE_OBSERVATION")
        literal={
            "schema_version":1,"scope":BOOK_SCOPE,
            "observed_on":"2026-10-12",
            "original_statistical_publisher":"U.S. Bureau of Economic Analysis",
            "official_first_party_grossoutput_xlsx_url":BEA_URL,
            "official_first_party_grossoutput_xlsx_sha256":BEA_ORIGINAL_SHA256,
            "official_first_party_grossoutput_xlsx_bytes":BEA_ORIGINAL_BYTES,
            "official_first_party_xlsx_sheet":"TGO103-A",
            "official_first_party_xlsx_reference_year":2017,
            "official_first_party_xlsx_row_number":52,
            "official_first_party_xlsx_industry_line":"44",
            "official_first_party_xlsx_industry_title":"Truck transportation",
            "official_first_party_xlsx_pub_date":"Data published September 30, 2026",
            "target_table_id":"16","target_naics":"484","target_frequency":"A",
            "target_years":list(YEARS),
            "original_bea_gross_output_primary_data_sha256":BEA_ORIGINAL_SHA256,
            "original_bea_numeric_source_authenticated_via_public_xlsx":True,
            "official_first_party_xlsx_contains_real_gross_output_values":True,
            "official_first_party_xlsx_and_dbnomics_values_reconcile":True,
            "official_first_party_numeric_xlsx_retrieval_independent_of_bea_api_key":True,
            "price_and_quantity_reconcile_to_nominal_within_0_02_pct":True,
            "official_first_party_xlsx_sheetnames":list(EXPECTED_SHEETS),
            "official_first_party_xlsx_units_by_sheet":{
                "TGO103-A":"INDEX_2017_EQUALS_100",
                "TGO104-A":"PRICE_INDEX_2017_EQUALS_100",
                "TGO105-A":"CURRENT_DOLLAR_MILLIONS",
                "TGO106-A":"2017_CHAINED_DOLLAR_MILLIONS",
            },
            "original_bea_api_authentication_satisfied":False,
            "original_bea_api_active_userid_available":False,
            "official_api_expected_error_code":"1",
            "official_api_expected_error_message":"Invalid Request - Invalid API UserId.",
            "official_bea_xlsx_and_dbnomics_values_reconcile":True,
        }
        for key,expected in literal.items():
            if type(b.get(key))!=type(expected) or b.get(key)!=expected:
                errors.append("M33_PRIMARY_SOURCE_IDENTITY_OR_MEASUREMENT_DRIFT:"+key)
        annual=b.get("official_first_party_xlsx_annual_source_values")
        if type(annual) is not dict or set(annual)!=set(SHEETS):
            errors.append("M33_ORIGINAL_BEA_4_ANNUAL_TABLES_MISSING")
        else:
            for sheet in SHEETS:
                expected={str(y):v for y,v in zip(YEARS,EXPECTED[sheet])}
                if annual.get(sheet)!=expected:
                    errors.append("M33_ORIGINAL_BEA_8_YEAR_TABLE_VALUES_CHANGED:"+sheet)
        if b.get("original_bea_gross_output_primary_year_values")!={
            str(y):v for y,v in zip(YEARS,EXPECTED["TGO103-A"])
        }:
            errors.append("M33_ORIGINAL_BEAS_GROSS_OUTPUT_484_VALUES_CHANGED")
        for name in FALSE_GUARDS:
            # One guard explicitly records a true prohibition, not a
            # permission: direct BEA file does not certify a recovery market.
            want=name=="direct_bea_file_does_not_certify_recoverable_freight_market"
            if b.get(name) is not want:
                errors.append("M33_FALSE_RECOVERY_OR_SOURCE_AUTHORIZATION:"+name)
        for field,v in SOURCE_NO_RECOVERY.items():
            if type(b.get(field))!=type(v) or b.get(field)!=v:
                errors.append("M33_CUSTOMER_INVOICE_OR_RECOVERY_VETO_REMOVED:"+field)
        note=b.get("source_caveat")
        if type(note) is not str or len(note)<900 or "eighteen" not in note.lower():
            errors.append("M33_MISSING_BEA_PUBLISHER_SOURCE_AND_CUSTOMER_LIMITS")
        return errors
    except (TypeError,ValueError,KeyError,AttributeError):
        return ["M33_BEA_SOURCE_BOOK_UNPARSABLE"]


def _year_header(ws):
    """Find real annual year headers, not hard-coded erroneous column guesses."""
    for row in ws.iter_rows(min_row=7,max_row=16,values_only=True):
        positions={}
        for col,value in enumerate(row,start=1):
            v=str(value).strip() if value is not None else ""
            if v in {str(y) for y in range(1997,2026)}:
                if v in positions:
                    raise ValueError("M33_DUPLICATE_BEA_ANNUAL_YEAR_HEADER")
                positions[v]=col
        if set(str(y) for y in YEARS).issubset(positions):
            cols=[positions[str(y)] for y in YEARS]
            if cols!=list(range(cols[0],cols[0]+len(YEARS))):
                raise ValueError("M33_NONCONTIGUOUS_BEA_ANNUAL_YEAR_COLUMNS")
            return {y:positions[str(y)] for y in YEARS}
    raise ValueError("M33_ORIGINAL_BEA_WORKBOOK_ANNUAL_YEAR_HEADERS_MISSING")


def inspect_xlsx_content(raw:bytes) -> dict:
    """A pure workbook validator; alone it does NOT prove BEA publisher origin."""
    from openpyxl import load_workbook
    if type(raw) is not bytes:
        raise ValueError("M33_ORIGINAL_XLSX_PAYLOAD_NOT_BYTES")
    try:
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=True)
    except Exception as exc:
        raise ValueError("M33_MALFORMED_BEA_ORIGINAL_XLSX") from exc
    try:
        if tuple(wb.sheetnames)!=EXPECTED_SHEETS:
            raise ValueError("M33_BEA_ORIGINAL_SOURCE_WORKBOOK_TABLE_SETS_CHANGED")
        tab=wb["Contents"]
        entries={
            str(row[1]):str(row[2])
            for row in tab.iter_rows(values_only=True)
            if len(row)>2 and row[1] is not None and row[2] is not None
        }
        if entries.get("TGO103")!=TITLES["TGO103-A"]:
            raise ValueError("M33_CONTENTS_GROSS_OUTPUT_CHAIN_QUANTITY_TABLE_MISSING")
        values={}
        source_year_columns={}
        for name in SHEETS:
            ws=wb[name]
            for row,val in (
                (1,TITLES[name]),
                (2,UNITS[name]),
                (3,"Annual data from 1997 to 2025"),
                (4,"Bureau of Economic Analysis"),
                (5,"Data published September 30, 2026"),
            ):
                if str(ws.cell(row,1).value)!=val:
                    raise ValueError("M33_BEA_ANNUAL_TABLE_HEADING_UNIT_OR_VINTAGE_CHANGED:"+name+":"+str(row))
            year_columns=_year_header(ws)
            if str(ws.cell(52,1).value)!="44" or str(ws.cell(52,2).value)!="Truck transportation":
                raise ValueError("M33_BEA_484_TRUCK_TRANSPORTATION_INDUSTRY_ROW_CHANGED:"+name)
            count=sum(
                1 for row in ws.iter_rows(min_row=14,max_row=90,min_col=2,max_col=2,values_only=True)
                if row[0]=="Truck transportation"
            )
            if count!=1:
                raise ValueError("M33_DUPLICATE_OR_MISSING_TRUCK_INDUSTRY_CLASSIFICATION:"+name)
            annual={}
            for year in YEARS:
                value=_decimal(ws.cell(52,year_columns[year]).value)
                if value!=Decimal(EXPECTED[name][year-2017]):
                    raise ValueError("M33_BEA_PRIMARY_ANNUAL_SOURCE_VALUE_DRIFT_REQUIRES_REVIEW:"+name+":"+str(year))
                annual[str(year)]=str(value)
            values[name]=annual
            source_year_columns[name]={str(y):year_columns[y] for y in YEARS}
        # Gross output price * quantity = nominal, with publication rounded
        # indexes and annual dollar amounts. A maximum relative error
        # of 20 ppm accommodates three-decimal source rounding, not actual
        # financial invoice reconciliation.
        max_err=Decimal("0")
        for year in YEARS:
            y=str(year)
            price=Decimal(values["TGO104-A"][y])
            quantity=Decimal(values["TGO103-A"][y])
            nominal=Decimal(values["TGO105-A"][y])
            reconstructed=(Decimal("367154")*price*quantity)/Decimal(10000)
            relative_err=abs(reconstructed/nominal-1)
            max_err=max(max_err,relative_err)
            if relative_err>Decimal("0.00002"):
                raise ValueError("M33_ORIGINAL_BEA_PRICE_QUANTITY_NOMINAL_IDENTITY_BROKEN:"+y)
            chained=Decimal(values["TGO106-A"][y])
            chain_expected=Decimal("367154")*quantity/Decimal(100)
            if abs(chain_expected-chained)>Decimal("6"):
                raise ValueError("M33_BEA_CHAIN_DOLLAR_VS_QUANTITY_SOURCE_DRIFT:"+y)
        return {
            "source_measure":"BEA_ORIGINAL_TGO103_ANNUAL_REAL_GROSS_OUTPUT_QUANTITY",
            "industry_naics":"484","industry":"Truck transportation",
            "source_rows_reproduced":4,
            "original_year_source_cells_reproduced":4*len(YEARS),
            "index_2017":values["TGO103-A"]["2017"],
            "index_2024":values["TGO103-A"]["2024"],
            "index_2017_to_2024_growth_pct":"8.05",
            "publisher_release_date":"2026-09-30",
            "source_annual_values":values,
            "source_year_excel_columns":source_year_columns,
            "max_price_times_quantity_vs_nominal_relative_error":str(max_err),
            "independent_customer_invoices_or_recovered_cash":0,
            "original_file_alone_not_new_statistically_independent_survey":True,
        }
    finally:
        wb.close()


def verify_original_publisher(*,as_of:date, raw:bytes)->dict:
    if type(raw) is not bytes or len(raw)!=BEA_ORIGINAL_BYTES:
        raise ValueError("M33_ORIGINAL_BEA_XLSX_BYTE_LENGTH_MISMATCH")
    if sha256(raw).hexdigest()!=BEA_ORIGINAL_SHA256:
        raise ValueError("M33_ORIGINAL_BEA_XLSX_SHA256_CHANGED")
    parsed=inspect_xlsx_content(raw)
    report=report_shell(as_of=as_of)
    report.update({
        "status":VERIFIED,
        "fresh_original_bea_publisher_xlsx_replayed_this_run":True,
        "fresh_original_bea_table16_year_cells_authenticated":8,
        "fresh_four_original_bea_tables_year_cells_authenticated":32,
        "direct_bea_original_numeric_source_authenticated":True,
        "fresh_original_bea_xlsx_sha256":sha256(raw).hexdigest(),
        "source_grade":"FIRST_PARTY_BEA_ORIGINAL_GROSSOUTPUT_XLSX_SHA256_VERIFIED",
        "original_bea_workbook_internal_reconciliation":parsed,
        "fresh_primary_source_semantic_receipt_sha256":canon_hash(parsed),
    })
    report["receipt_sha256"]=canon_hash({
        k:v for k,v in report.items() if k!="receipt_sha256"
    })
    return report


def _baseline(*,as_of:date)->tuple[dict,dict,dict]:
    catalog=_book()
    problems=validate_book(catalog,as_of=as_of)
    if problems:
        raise ValueError("M33_BEAS_SOURCE_BOOK_INVALID: "+";".join(problems))
    m32=m32_redistributor_report(as_of=as_of)
    m31=m31_sectoral_output_report(as_of=as_of)
    if (m32["status"]!=M32_OFFLINE or
        m32["bea_real_gross_output_2017_2024_change_pct"]!="8.05" or
        m32["direct_original_bea_table16_numeric_source_bytes_authenticated"] is not False or
        m32["human_pilot_18_independent_requirements_unmet"] is not True or
        m31["status"]!=M31_OFFLINE or
        m31["aggregate_naics484_real_sectoral_output_2017_2024_pct"]!="4.02"):
        raise ValueError("M33_PREVIOUS_BEA_SECONDARY_OR_BLS_ORIGINAL_SOURCE_LINEAGE_DRIFT")
    return catalog,m32,m31


def report_shell(*,as_of:date)->dict:
    catalog,m32,m31=_baseline(as_of=as_of)
    price=catalog["official_first_party_xlsx_annual_source_values"]["TGO104-A"]
    quantity=catalog["official_first_party_xlsx_annual_source_values"]["TGO103-A"]
    nominal=catalog["official_first_party_xlsx_annual_source_values"]["TGO105-A"]
    chain=catalog["official_first_party_xlsx_annual_source_values"]["TGO106-A"]
    rows=[
        {
            "year":year,
            "first_party_bea_real_gross_output_quantity_index":quantity[str(year)],
            "first_party_bea_gross_output_price_index":price[str(year)],
            "first_party_bea_current_dollar_gross_output_million_usd":nominal[str(year)],
            "first_party_bea_real_chained_2017_gross_output_million_usd":chain[str(year)],
            "reproduced_from_original_xlsx_in_this_report":False,
            "is_customer_invoice_or_recoverable_cash":False,
        }
        for year in YEARS
    ]
    if [x["first_party_bea_real_gross_output_quantity_index"]
        for x in rows]!=[
        y["chain_type_real_gross_output_index"] for y in m32["bea_2017_2024_year_index"]
    ]:
        raise ValueError("M33_PRIMARY_BEA_AND_SECONDARY_DBNOMICS_ANNUAL_484_VALUES_CONFLICT")
    result={
        "schema_version":1,"scope":SCOPE,"as_of":as_of.isoformat(),
        "lab":14,"new_lab_created":False,
        "status":OFFLINE,
        "source_url":BEA_URL,
        "source_publisher":"U.S. Bureau of Economic Analysis",
        "original_bea_file_expected_sha256":BEA_ORIGINAL_SHA256,
        "original_bea_file_expected_bytes":BEA_ORIGINAL_BYTES,
        "fresh_original_bea_publisher_xlsx_replayed_this_run":False,
        "fresh_original_bea_table16_year_cells_authenticated":0,
        "fresh_four_original_bea_tables_year_cells_authenticated":0,
        "direct_bea_original_numeric_source_authenticated":False,
        "fresh_original_bea_xlsx_sha256":None,
        "fresh_primary_source_semantic_receipt_sha256":None,
        "source_grade":"OFFLINE_PRIOR_DIRECT_BEA_PUBLICATION_SOURCE_MANIFEST_NO_NEW_DOWNLOAD",
        "original_bea_workbook_internal_reconciliation":None,
        "official_bea_api_userid_denial_remains_true_historical_event":True,
        "original_bea_file_download_does_not_require_bea_api_registration":True,
        "previous_m32_secondary_source_only_remains_historical":True,
        "bea_484_2017_2024_gross_output_annual_observations":rows,
        "original_bea_gross_output_2017_to_2024_pct":"8.05",
        "bls_real_sectoral_output_2017_to_2024_pct":"4.02",
        "illustrative_census_bls_ppi_revenue_growth_pct":"1.18",
        "bea_price_index_2017_to_2024_pct":"37.49",
        "bea_nominal_gross_output_2017_million_usd":nominal["2017"],
        "bea_nominal_gross_output_2024_million_usd":nominal["2024"],
        "bea_nominal_gross_output_2017_to_2024_pct":"48.56",
        "parent_naics484_includes_both_truckload_and_ltl":True,
        "economic_measure_discrepancies_are_not_forecast_errors":True,
        "publisher_original_file_is_not_independent_paid_customer_audit":True,
        "source_book_receipt_sha256":canon_hash(catalog),
        "prior_m32_secondary_bea_receipt_sha256":m32["receipt_sha256"],
        "prior_m31_original_bls_receipt_sha256":m31["receipt_sha256"],
        "paid_customer_freight_invoice_rows":0,
        "verified_carrier_invoice_overcharges":0,
        "actual_retally_recovery_cash_usd":None,
        "recoverable_freight_market_tam_usd":None,
        "buyers_with_authorized_verified_demand":0,
        "external_recovery_performance_claims_authorized":False,
        "private_customer_or_carrier_data_ingested":False,
        "human_pilot_18_independent_requirements_unmet":True,
        "scheduled_tasks_created_or_modified":False,
        "caveat":(
            "BEA itself publishes a GrossOutput.xlsx industry workbook on "
            "apps.bea.gov. This source contains NAICS 484 Truck transportation "
            "annual TGO103-A Chain-Type Quantity Indexes for Gross Output "
            "by Industry (2017=100) and matching official gross-output "
            "price, current-dollar and chained-dollar tables. Only a fresh "
            "byte-length and SHA-256 verified first-party BEA publisher "
            "download can mark original BEA numeric source reproduced in "
            "this execution; offline reports cannot claim fresh source proof. "
            "Original BEA current-dollar gross output and sectoral quantity "
            "indexes measure economy-wide industry production, NOT paid "
            "customer freight invoices or serviceable LTL freight audit "
            "spending. Original BEA 8.05% real gross output differs from "
            "BLS 4.02% real sectoral output and an illustrative Census/PPI "
            "1.18% proxy. Their differences are not statistically established "
            "model forecast errors, cargo tonnage, billing discrepancies, "
            "legally collectable refunds, or evidence of customer demand. "
            "BEA price-times-quantity within a single BEA release is a "
            "consistency check, not a second independent survey."
        ),
    }
    result["receipt_sha256"]=canon_hash(result)
    return result


def blocked(*,as_of:date)->dict:
    result=report_shell(as_of=as_of)
    result["status"]=BLOCKED
    result["receipt_sha256"]=canon_hash({
        k:v for k,v in result.items() if k!="receipt_sha256"
    })
    return result


def fetch_original_bea()->bytes:
    req=Request(BEA_URL,headers={"User-Agent":"RETALLY-Lab14-Original-BEA-Gross-Output-Source/1.0"})
    with urlopen(req,timeout=35) as response:
        if response.geturl()!=BEA_URL:
            raise ValueError("M33_FIRST_PARTY_BEA_PUBLISHER_REDIRECTED")
        raw=response.read(MAX_BYTES+1)
    if len(raw)>MAX_BYTES:
        raise ValueError("M33_ORIGINAL_BEA_XLSX_EXCEEDED_MAXIMUM_ALLOWED_BYTES")
    return raw


def render_html(r:dict)->str:
    rows="".join(
        "<tr><th>"+str(x["year"])+"</th><td>"+
        escape(x["first_party_bea_real_gross_output_quantity_index"])+
        "</td><td>"+
        escape(x["first_party_bea_gross_output_price_index"])+
        "</td><td>"+
        escape(x["first_party_bea_current_dollar_gross_output_million_usd"])+
        "</td><td>"+
        escape(x["first_party_bea_real_chained_2017_gross_output_million_usd"])+
        "</td></tr>"
        for x in r["bea_484_2017_2024_gross_output_annual_observations"]
    )
    return (
        '<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<title>RETALLY | First-party BEA original gross output verified</title>'
        '<style>body{font:15px/1.5 system-ui;background:#f6f9f7;color:#173a2b;'
        'margin:0}main{max-width:1100px;margin:36px auto;padding:18px}'
        'table{border-collapse:collapse;width:100%;background:#fff}'
        'td,th{border-bottom:1px solid #d4dfd6;padding:11px;text-align:left}'
        'small{color:#557063}</style></head><body><main>'
        '<small>RETALLY / ORIGINAL RESEARCH LAB 14 / MISSION 33</small>'
        '<h1>Original BEA gross output, trucking NAICS 484</h1>'
        '<p>Publisher numeric file source status: <strong>'+
        escape(r["status"])+'</strong>. Fresh original first-party BEA '
        'source years authenticated this run: <strong>'+
        str(r["fresh_original_bea_table16_year_cells_authenticated"])+
        ' of 8</strong>. Offline analysis shows prior pinned values and '
        'MUST NOT claim a new publisher download.</p>'
        '<table><thead><tr><th>Year</th><th>Real gross-output quantity '
        '(index 2017=100)</th><th>Gross-output price index</th>'
        '<th>Nominal gross output ($million)</th><th>Real gross output '
        '($million chained 2017)</th></tr></thead><tbody>'+rows+
        '</tbody></table><p>BEA real gross-output quantity change 2017–2024: '
        '<strong>'+escape(r["original_bea_gross_output_2017_to_2024_pct"])+
        '%</strong>. BLS sectoral output: '+
        escape(r["bls_real_sectoral_output_2017_to_2024_pct"])+
        '%. Illustrative Census revenue/PPI: '+
        escape(r["illustrative_census_bls_ppi_revenue_growth_pct"])+
        '%. <strong>These measures are not additive, interchangeable, '
        'independent forecasts or customer freight audits.</strong></p>'
        '<p>'+escape(r["caveat"])+'</p>'
        '<p><a href="'+BEA_URL+'">BEA original GrossOutput.xlsx</a></p>'
        '</main></body></html>'
    )


def main()->None:
    p=argparse.ArgumentParser(description="BEA original XLSX Table16 source replay, research only")
    p.add_argument("--out",type=Path,required=True)
    p.add_argument("--as-of-date",default=date.today().isoformat())
    mode=p.add_mutually_exclusive_group(required=True)
    mode.add_argument("--offline",action="store_true")
    mode.add_argument("--fetch-original",action="store_true")
    args=p.parse_args()
    as_of=date.fromisoformat(args.as_of_date)
    if args.fetch_original:
        try: raw=fetch_original_bea()
        except (OSError,TimeoutError):
            result=blocked(as_of=as_of)
        else:
            result=verify_original_publisher(as_of=as_of,raw=raw)
    else:
        result=report_shell(as_of=as_of)
    args.out.mkdir(parents=True,exist_ok=True)
    (args.out/"bea_original_gross_output.json").write_text(
        json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8"
    )
    (args.out/"bea_original_gross_output.html").write_text(
        render_html(result),encoding="utf-8"
    )
    print(json.dumps({
        "source_status":result["status"],
        "new_first_party_bea_table16_years":result["fresh_original_bea_table16_year_cells_authenticated"],
        "new_first_party_bea_four_table_cells":result["fresh_four_original_bea_tables_year_cells_authenticated"],
        "primary_numeric_source_authenticated_in_this_run":result["direct_bea_original_numeric_source_authenticated"],
        "real_customer_invoice_overcharges":0,
        "receipt_sha256":result["receipt_sha256"],
    },sort_keys=True))


if __name__=="__main__":
    main()
