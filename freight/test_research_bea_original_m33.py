"""Mission 33 adversarial original BEA XLSX source controls.

Synthetic XLSX workbooks in these tests DO NOT authenticate original BEA.
Only a live GET from apps.bea.gov with full pinned byte count, SHA256 and
32 strict publisher cells can establish direct original-source proof.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from freight.research_evidence_gate import canon_hash
from freight import research_bea_original_m33 as m33

TODAY=date(2026,10,12)


def catalog():
    return json.loads(m33.BOOK.read_text(encoding="utf-8"))


def synthetic_xlsx():
    wb=Workbook()
    wb.active.title="Contents"
    for sh in m33.EXPECTED_SHEETS[1:]:
        wb.create_sheet(sh)
    sheet=wb["Contents"]
    sheet["B6"]="TGO103"
    sheet["C6"]=m33.TITLES["TGO103-A"]
    for name in m33.SHEETS:
        ws=wb[name]
        for row,val in (
            (1,m33.TITLES[name]),
            (2,m33.UNITS[name]),
            (3,"Annual data from 1997 to 2025"),
            (4,"Bureau of Economic Analysis"),
            (5,"Data published September 30, 2026"),
        ):
            ws.cell(row=row,column=1,value=val)
        for year in range(1997,2026):
            ws.cell(row=8,column=year-1997+4,value=year)
        ws.cell(52,1,44)
        ws.cell(52,2,"Truck transportation")
        for y,val in zip(m33.YEARS,m33.EXPECTED[name]):
            number=Decimal(val)
            ws.cell(52,y-1997+4,float(number) if "." in val else int(number))
    output=BytesIO()
    wb.save(output)
    return output.getvalue()


def mutate_fixture(func):
    from openpyxl import load_workbook
    w=load_workbook(BytesIO(synthetic_xlsx()))
    func(w)
    output=BytesIO()
    w.save(output)
    return output.getvalue()


class OriginalBEAXLSXTests(unittest.TestCase):
    def test_original_publisher_manifest_git_blob_and_unit_controls(self):
        self.assertEqual(m33._git_blob(m33.BOOK.read_bytes()),m33.BOOK_SHA1)
        self.assertEqual(m33.validate_book(catalog(),as_of=TODAY),[])
        self.assertEqual(catalog()["official_first_party_grossoutput_xlsx_url"],m33.BEA_URL)
        self.assertEqual(catalog()["official_first_party_grossoutput_xlsx_bytes"],506819)
        self.assertEqual(catalog()["official_first_party_grossoutput_xlsx_sha256"],
                         m33.BEA_ORIGINAL_SHA256)

    def test_source_is_bea_first_party_not_fred_or_dbnomics(self):
        r=m33.report_shell(as_of=TODAY)
        self.assertEqual(r["source_publisher"],"U.S. Bureau of Economic Analysis")
        self.assertEqual(r["source_url"],"https://apps.bea.gov/industry/Release/XLS/GDPxInd/GrossOutput.xlsx")
        self.assertTrue(r["previous_m32_secondary_source_only_remains_historical"])
        self.assertTrue(r["official_bea_api_userid_denial_remains_true_historical_event"])

    def test_original_grossoutput_workbook_four_tables_and_reference_year(self):
        b=catalog()
        self.assertEqual(b["official_first_party_xlsx_sheet"],"TGO103-A")
        self.assertEqual(b["official_first_party_xlsx_pub_date"],
                         "Data published September 30, 2026")
        self.assertEqual(b["official_first_party_xlsx_industry_line"],"44")
        self.assertEqual(b["official_first_party_xlsx_row_number"],52)
        self.assertEqual(b["official_first_party_xlsx_sheetnames"],list(m33.EXPECTED_SHEETS))
        self.assertEqual(set(b["official_first_party_xlsx_annual_source_values"]),set(m33.SHEETS))

    def test_synthetic_four_table_xlsx_parser_reads_all_32_year_cells(self):
        parsed=m33.inspect_xlsx_content(synthetic_xlsx())
        self.assertEqual(parsed["original_year_source_cells_reproduced"],32)
        self.assertEqual(parsed["source_rows_reproduced"],4)
        self.assertEqual(parsed["publisher_release_date"],"2026-09-30")
        self.assertEqual(Decimal(parsed["index_2017"]),Decimal("100"))
        self.assertEqual(parsed["index_2024"],"108.047")
        self.assertEqual(parsed["index_2017_to_2024_growth_pct"],"8.05")
        self.assertEqual(parsed["source_year_excel_columns"]["TGO103-A"]["2017"],24)
        self.assertEqual(parsed["source_year_excel_columns"]["TGO103-A"]["2024"],31)

    def test_quantity_series_reconciles_secondary_dbnomics_values(self):
        r=m33.report_shell(as_of=TODAY)
        indices=[x["first_party_bea_real_gross_output_quantity_index"]
                 for x in r["bea_484_2017_2024_gross_output_annual_observations"]]
        self.assertEqual(indices,list(m33.EXPECTED["TGO103-A"]))
        self.assertEqual(r["original_bea_gross_output_2017_to_2024_pct"],"8.05")
        self.assertEqual(r["bls_real_sectoral_output_2017_to_2024_pct"],"4.02")
        self.assertEqual(r["illustrative_census_bls_ppi_revenue_growth_pct"],"1.18")

    def test_bea_nominal_and_chain_output_source_units_not_mixed(self):
        r=m33.report_shell(as_of=TODAY)
        rows=r["bea_484_2017_2024_gross_output_annual_observations"]
        self.assertEqual(rows[0]["first_party_bea_current_dollar_gross_output_million_usd"],"367154")
        self.assertEqual(rows[-1]["first_party_bea_current_dollar_gross_output_million_usd"],"545435")
        self.assertEqual(rows[-1]["first_party_bea_real_chained_2017_gross_output_million_usd"],"396699")
        self.assertEqual(rows[-1]["first_party_bea_gross_output_price_index"],"137.494")
        self.assertEqual(r["bea_price_index_2017_to_2024_pct"],"37.49")
        self.assertEqual(r["bea_nominal_gross_output_2017_to_2024_pct"],"48.56")

    def test_recomputed_price_and_quantity_reconcile_to_nominal_gross_output(self):
        parsed=m33.inspect_xlsx_content(synthetic_xlsx())
        self.assertLessEqual(
            Decimal(parsed["max_price_times_quantity_vs_nominal_relative_error"]),
            Decimal("0.00002")
        )
        rows=parsed["source_annual_values"]
        self.assertEqual(rows["TGO103-A"]["2024"],"108.047")
        self.assertEqual(rows["TGO104-A"]["2024"],"137.494")
        self.assertEqual(rows["TGO105-A"]["2024"],"545435")
        self.assertEqual(rows["TGO106-A"]["2024"],"396699")

    def test_offline_cannot_claim_new_fresh_bea_source(self):
        r=m33.report_shell(as_of=TODAY)
        self.assertEqual(r["status"],m33.OFFLINE)
        self.assertFalse(r["fresh_original_bea_publisher_xlsx_replayed_this_run"])
        self.assertFalse(r["direct_bea_original_numeric_source_authenticated"])
        self.assertEqual(r["fresh_original_bea_table16_year_cells_authenticated"],0)
        self.assertEqual(r["fresh_four_original_bea_tables_year_cells_authenticated"],0)
        self.assertIsNone(r["fresh_original_bea_xlsx_sha256"])
        self.assertIsNone(r["original_bea_workbook_internal_reconciliation"])

    def test_wrong_first_party_xlsx_byte_length_rejected(self):
        with self.assertRaisesRegex(ValueError,"M33_ORIGINAL_BEA_XLSX_BYTE_LENGTH_MISMATCH"):
            m33.verify_original_publisher(as_of=TODAY,raw=synthetic_xlsx())

    def test_exact_sized_but_replaced_original_bea_xlsx_digest_rejected(self):
        with self.assertRaisesRegex(ValueError,"M33_ORIGINAL_BEA_XLSX_SHA256_CHANGED"):
            m33.verify_original_publisher(as_of=TODAY,raw=b"x"*m33.BEA_ORIGINAL_BYTES)

    def test_wrong_workbook_sheet_names_rejected(self):
        raw=mutate_fixture(lambda w: w["TGO103-A"].__setattr__("title","TGO999-A"))
        with self.assertRaisesRegex(ValueError,"M33_BEA_ORIGINAL_SOURCE_WORKBOOK_TABLE_SETS_CHANGED"):
            m33.inspect_xlsx_content(raw)

    def test_wrong_or_missing_contents_table_identity_rejected(self):
        raw=mutate_fixture(lambda w: w["Contents"].__setitem__("C6","Customer refunds"))
        with self.assertRaisesRegex(ValueError,"M33_CONTENTS_GROSS_OUTPUT_CHAIN_QUANTITY_TABLE_MISSING"):
            m33.inspect_xlsx_content(raw)

    def test_wrong_measure_title_is_not_gross_output(self):
        for tab in m33.SHEETS:
            raw=mutate_fixture(lambda w, name=tab: w[name].__setitem__("A1","GDP Value Added by Industry"))
            with self.assertRaisesRegex(ValueError,
                                        "M33_BEA_ANNUAL_TABLE_HEADING_UNIT_OR_VINTAGE_CHANGED",msg=tab):
                m33.inspect_xlsx_content(raw)

    def test_wrong_source_price_index_units_are_not_quantity_index(self):
        raw=mutate_fixture(lambda w: w["TGO103-A"].__setitem__("A2","[Percent change]"))
        with self.assertRaisesRegex(ValueError,"M33_BEA_ANNUAL_TABLE_HEADING_UNIT_OR_VINTAGE_CHANGED"):
            m33.inspect_xlsx_content(raw)

    def test_older_bea_release_vintage_must_fail(self):
        raw=mutate_fixture(lambda w: w["TGO103-A"].__setitem__("A5","Data published September 30, 2025"))
        with self.assertRaisesRegex(ValueError,"M33_BEA_ANNUAL_TABLE_HEADING_UNIT_OR_VINTAGE_CHANGED"):
            m33.inspect_xlsx_content(raw)

    def test_wrong_484_industry_label_or_line_fails(self):
        for cell,val in (("A52",43),("B52","Water transportation")):
            raw=mutate_fixture(lambda w, c=cell,v=val:w["TGO103-A"].__setitem__(c,v))
            with self.assertRaisesRegex(ValueError,
                                        "M33_BEA_484_TRUCK_TRANSPORTATION_INDUSTRY_ROW_CHANGED",msg=cell):
                m33.inspect_xlsx_content(raw)

    def test_two_duplicate_truck_transportation_rows_rejected(self):
        raw=mutate_fixture(lambda w:w["TGO103-A"].__setitem__("B51","Truck transportation"))
        with self.assertRaisesRegex(ValueError,
                                    "M33_DUPLICATE_OR_MISSING_TRUCK_INDUSTRY_CLASSIFICATION"):
            m33.inspect_xlsx_content(raw)

    def test_off_by_one_year_header_fails(self):
        raw=mutate_fixture(lambda w: w["TGO103-A"].__setitem__("Y8",2035))
        with self.assertRaisesRegex(ValueError,
                                    "M33_ORIGINAL_BEA_WORKBOOK_ANNUAL_YEAR_HEADERS_MISSING"):
            m33.inspect_xlsx_content(raw)

    def test_moved_year_value_fails_even_if_all_eight_have_data(self):
        for y in (2017,2020,2024):
            col=y-1997+4
            raw=mutate_fixture(lambda w,c=col: w["TGO103-A"].cell(52,c).__setattr__("value",107.7))
            with self.assertRaisesRegex(ValueError,
                                        "M33_BEA_PRIMARY_ANNUAL_SOURCE_VALUE_DRIFT_REQUIRES_REVIEW",msg=str(y)):
                m33.inspect_xlsx_content(raw)

    def test_one_cell_wrong_in_each_of_four_tables_fails_closed(self):
        for name in m33.SHEETS:
            raw=mutate_fixture(lambda w,n=name:w[n].__setitem__("AE52",999.99))
            with self.assertRaisesRegex(ValueError,
                                        "M33_BEA_PRIMARY_ANNUAL_SOURCE_VALUE_DRIFT_REQUIRES_REVIEW",msg=name):
                m33.inspect_xlsx_content(raw)

    def test_non_numeric_suppressed_and_boolean_cells_rejected(self):
        for value in ("S","NaN",False,None,-1,0,{},1e20):
            with self.assertRaises(ValueError,msg=repr(value)):
                m33._decimal(value)

    def test_gross_output_not_equivalent_to_bls_sectoral_or_recovery_market(self):
        r=m33.report_shell(as_of=TODAY)
        self.assertTrue(r["economic_measure_discrepancies_are_not_forecast_errors"])
        self.assertTrue(r["parent_naics484_includes_both_truckload_and_ltl"])
        self.assertTrue(r["publisher_original_file_is_not_independent_paid_customer_audit"])
        self.assertEqual(r["paid_customer_freight_invoice_rows"],0)
        self.assertIsNone(r["recoverable_freight_market_tam_usd"])
        self.assertTrue(r["human_pilot_18_independent_requirements_unmet"])

    def test_m33_catalog_forged_price_or_quantity_values_blocked(self):
        for sheet in m33.SHEETS:
            b=catalog()
            b["official_first_party_xlsx_annual_source_values"][sheet]["2024"]="9999"
            self.assertTrue(m33.validate_book(b,as_of=TODAY),sheet)

    def test_m33_catalog_forged_source_digest_or_url_blocked(self):
        for key,val in (
            ("official_first_party_grossoutput_xlsx_url","https://attacker.invalid/file.xlsx"),
            ("official_first_party_grossoutput_xlsx_sha256","0"*64),
            ("official_first_party_grossoutput_xlsx_bytes",12345),
            ("official_first_party_xlsx_pub_date","Data published September 30, 2025"),
            ("original_bea_gross_output_primary_data_sha256",None),
        ):
            b=catalog()
            b[key]=val
            self.assertTrue(m33.validate_book(b,as_of=TODAY),key)

    def test_api_key_was_not_created_and_old_api_diagnostic_still_true(self):
        b=catalog()
        self.assertFalse(b["original_bea_api_authentication_satisfied"])
        self.assertFalse(b["original_bea_api_active_userid_available"])
        self.assertTrue(b["official_first_party_numeric_xlsx_retrieval_independent_of_bea_api_key"])
        self.assertEqual(b["official_api_expected_error_code"],"1")
        self.assertTrue(m33.report_shell(as_of=TODAY)["official_bea_api_userid_denial_remains_true_historical_event"])

    def test_source_book_does_not_allow_customer_invoices_or_refunds(self):
        for key in (
            "actual_real_customer_invoice_rows",
            "actual_independently_confirmed_invoice_overcharges",
            "real_refunds_usd",
            "recoverable_market_tam_usd",
            "actual_buyer_demand_established",
            "eighteen_independent_customer_pilot_approvals_satisfied",
            "scheduled_tasks_modified",
        ):
            b=catalog()
            b[key]=100
            self.assertTrue(m33.validate_book(b,as_of=TODAY),key)

    def test_source_observation_time_travel_rejected(self):
        self.assertTrue(m33.validate_book(catalog(),as_of=date(2026,10,11)))
        with self.assertRaisesRegex(ValueError,"M33_BEAS_SOURCE_BOOK_INVALID"):
            m33.report_shell(as_of=date(2026,10,11))

    def test_fetch_fixed_bea_original_url_and_no_registration(self):
        self.assertEqual(m33.BEA_URL,
                         "https://apps.bea.gov/industry/Release/XLS/GDPxInd/GrossOutput.xlsx")
        self.assertNotIn("api",m33.BEA_URL.lower())
        self.assertNotIn("UserID",m33.BEA_URL)

    def test_network_timeout_creates_blocked_not_source_attestation(self):
        r=m33.blocked(as_of=TODAY)
        self.assertEqual(r["status"],m33.BLOCKED)
        self.assertFalse(r["fresh_original_bea_publisher_xlsx_replayed_this_run"])
        self.assertEqual(r["fresh_original_bea_table16_year_cells_authenticated"],0)

    def test_cli_network_timeout_writes_distinct_blocked_receipt(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("sys.argv",["m33","--out",directory,"--as-of-date","2026-10-12",
                                   "--fetch-original"]):
                with patch.object(m33,"fetch_original_bea",side_effect=TimeoutError("test")):
                    m33.main()
            report=json.loads((Path(directory)/"bea_original_gross_output.json").read_text())
            self.assertEqual(report["status"],m33.BLOCKED)
            self.assertFalse(report["direct_bea_original_numeric_source_authenticated"])

    def test_changed_original_file_fails_not_network_block(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("sys.argv",["m33","--out",directory,"--as-of-date","2026-10-12",
                                   "--fetch-original"]):
                with patch.object(m33,"fetch_original_bea",return_value=b"tampered"):
                    with self.assertRaisesRegex(ValueError,"M33_ORIGINAL_BEA_XLSX_BYTE_LENGTH_MISMATCH"):
                        m33.main()

    def test_report_html_escapes_untrusted_caveat_and_preserves_caveats(self):
        report=m33.report_shell(as_of=TODAY)
        report["caveat"]="<script>misstate customer overcharges</script>"
        html=m33.render_html(report)
        self.assertIn("&lt;script&gt;",html)
        self.assertNotIn("<script>",html)
        self.assertIn("2017",html)
        self.assertIn("2024",html)
        self.assertIn(m33.OFFLINE,html)

    def test_offline_receipt_hash_determinism(self):
        r=m33.report_shell(as_of=TODAY)
        self.assertEqual(r["receipt_sha256"],canon_hash({
            k:v for k,v in r.items() if k!="receipt_sha256"
        }))


if __name__=="__main__":
    unittest.main()
