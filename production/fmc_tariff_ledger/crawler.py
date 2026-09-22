#!/usr/bin/env python3
"""FMC tariff directory crawler and versioned rule ledger.

Goals:
- Enumerate current + historical FMC tariff locations for VOCC, OTI/NVOCC,
  MTO, and conference records.
- Snapshot every discoverable tariff/rate/rule artifact as immutable bytes.
- Hash every snapshot with SHA-256.
- Extract effective dates, source versions, monetary amounts, free-time
  quantities, and rule terms into a carrier x rule x effective-date ledger.
- Fail visibly: blocked/auth/JS-only portals are recorded in crawl_errors.

This tool intentionally does not claim that a crawled page is a governing
contract. The ledger stores evidence and provenance; downstream FreightRecovery
logic decides applicability using the customer's agreement hierarchy.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import threading
import time
import urllib.parse
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Iterator

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


FMC_BASE = "https://www2.fmc.gov/FMC1Users/scripts/ExtReports.asp"
DIRECTORY_REPORTS = (
    ("vocc", True, f"{FMC_BASE}?tariffClass=vocc"),
    ("vocc", False, f"{FMC_BASE}?inactive=yes&tariffClass=vocc"),
    ("oti", True, f"{FMC_BASE}?tariffClass=oti"),
    ("oti", False, f"{FMC_BASE}?inactive=yes&tariffClass=oti"),
    ("mto", True, f"{FMC_BASE}?tariffClass=mto"),
    ("mto", False, f"{FMC_BASE}?inactive=yes&tariffClass=mto"),
    ("conference", True, f"{FMC_BASE}?tariffClass=conf"),
    ("conference", False, f"{FMC_BASE}?inactive=yes&tariffClass=conf"),
)

USER_AGENT = (
    "FMC-Tariff-Rule-Ledger/1.0 "
    "(public-tariff archival research; contact via repository)"
)

TARIFFISH = re.compile(
    r"(tariff|rate|charge|schedule|rule|demurr|detention|storage|wharf|dockage|"
    r"terminal|handling|fuel|surcharge|chassis|congestion|security|reefer|"
    r"hazard|local.?charge|free.?time|service.?contract|fmc)",
    re.I,
)
SUPPORTED_EXT = {
    ".html", ".htm", ".txt", ".csv", ".tsv", ".json", ".xml",
    ".pdf", ".xlsx", ".xlsm",
}
SKIP_EXT = {
    ".jpg", ".jpeg", ".png", ".gif", ".svg", ".ico", ".webp",
    ".mp4", ".mov", ".avi", ".mp3", ".zip", ".gz", ".tar", ".7z",
    ".woff", ".woff2", ".ttf", ".css", ".js",
}

RULE_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("demurrage", re.compile(r"\bdemurrage\b", re.I)),
    ("detention", re.compile(r"\bdetention\b", re.I)),
    ("free_time", re.compile(r"\bfree\s*time\b|\bfree\s+days?\b", re.I)),
    ("storage", re.compile(r"\bstorage\b", re.I)),
    ("terminal_handling", re.compile(r"terminal handling|\bTHC\b", re.I)),
    ("wharfage", re.compile(r"\bwharfage\b", re.I)),
    ("dockage", re.compile(r"\bdockage\b", re.I)),
    ("documentation", re.compile(r"documentation|document fee|B/?L fee|bill of lading fee", re.I)),
    ("fuel_surcharge", re.compile(r"fuel surcharge|\bBAF\b|bunker adjustment", re.I)),
    ("chassis", re.compile(r"\bchassis\b", re.I)),
    ("reefer", re.compile(r"reefer|refrigerat|genset", re.I)),
    ("hazardous", re.compile(r"hazardous|hazmat|dangerous goods|\bDG\b", re.I)),
    ("congestion", re.compile(r"congestion surcharge|port congestion", re.I)),
    ("security", re.compile(r"security fee|ISPS", re.I)),
    ("seal", re.compile(r"seal fee|container seal", re.I)),
    ("equipment", re.compile(r"equipment imbalance|equipment fee", re.I)),
    ("inland_haulage", re.compile(r"inland haulage|merchant haulage|carrier haulage", re.I)),
    ("cfs", re.compile(r"\bCFS\b|container freight station", re.I)),
    ("minimum_charge", re.compile(r"minimum charge|min(?:imum)?\. charge", re.I)),
    ("late_fee", re.compile(r"late fee|late charge", re.I)),
    ("general_rate", re.compile(r"ocean freight|base rate|freight rate|rate per", re.I)),
    ("effective_rule", re.compile(r"effective date|effective from|revised|revision|supplement", re.I)),
    ("applicability_date", re.compile(r"(?:rates?|charges?|rules?)[\s\S]{0,160}(?:in effect|effective)[\s\S]{0,160}(?:cargo|shipment)[\s\S]{0,100}received|date[\s\S]{0,100}(?:cargo|shipment)[\s\S]{0,100}received", re.I)),
    ("pass_through", re.compile(r"pass(?:ed)?[ -]?through|pass-through|cross-reference|without markup|not be marked up|no markup|at cost", re.I)),
]

MONEY_RE = re.compile(
    r"(?P<prefix>US\$|USD|\$)\s*"
    r"(?P<amount>\d{1,3}(?:,\d{3})*(?:\.\d{1,4})?|\d+(?:\.\d{1,4})?)"
    r"(?:\s*(?P<unit>/\s*[A-Za-z0-9'\-]+|per\s+(?:calendar\s+)?"
    r"(?:day|hour|container|unit|shipment|bill|document|B/?L|TEU|FEU|ton|"
    r"metric\s+ton|1000\s+kg|CBM)))?",
    re.I,
)
PERCENT_RE = re.compile(
    r"(?P<amount>\d+(?:\.\d+)?)\s*(?:%|percent)"
    r"(?:\s*(?:of\s+)?(?P<unit>[A-Za-z][A-Za-z0-9 \-/]{0,40}))?",
    re.I,
)
FREE_DAYS_RE = re.compile(
    r"(?:free\s*time|free\s+days?)[^\n.;:]{0,100}?"
    r"(?P<qty>\d{1,3})\s*(?:calendar\s+|working\s+|business\s+)?days?",
    re.I,
)
DATE_PATTERNS = [
    re.compile(
        r"(?:effective(?:\s+dates?)?|effective\s+from|effective\s+as\s+of|"
        r"revised|revision\s+date|valid\s+from|commencing)\s*[:\-]?\s*"
        r"(?P<date>\d{4}-\d{1,2}-\d{1,2}|\d{8}|"
        r"\d{1,2}/\d{1,2}/\d{2,4}|\d{1,2}-[A-Za-z]{3}-\d{4}|"
        r"\d{1,2}[A-Za-z]{3}\d{4}|"
        r"[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4}|"
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})",
        re.I,
    ),
]
END_DATE_PATTERNS = [
    re.compile(
        r"(?:expires?|expire\s+date|expiration(?:\s+date)?|"
        r"valid\s+(?:through|until)|effective\s+through)\s*[:\-]?\s*"
        r"(?P<date>\d{4}-\d{1,2}-\d{1,2}|\d{8}|"
        r"\d{1,2}/\d{1,2}/\d{2,4}|\d{1,2}-[A-Za-z]{3}-\d{4}|"
        r"\d{1,2}[A-Za-z]{3}\d{4}|"
        r"[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4}|"
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})",
        re.I,
    ),
]
VERSION_PATTERNS = [
    re.compile(r"tariff\s+(?:no\.?|number)\s*[:#]?\s*([A-Za-z0-9.\-_]+)", re.I),
    re.compile(r"(?:revision|rev\.)\s*[:#]?\s*([A-Za-z0-9.\-_]+)", re.I),
    re.compile(r"supplement\s+(?:no\.?|number)?\s*[:#]?\s*([A-Za-z0-9.\-_]+)", re.I),
]

PORTAL_DOMAINS = {
    "rates.descartes.com", "dpiusa.com", "www.dpiusa.com", "etmrates.com",
    "www.etmrates.com", "boterates.com", "www.boterates.com", "ratewave.com",
    "www.ratewave.com", "dmstradeservices.com", "www.dmstradeservices.com",
    "aptariffs.com", "www.aptariffs.com", "tariffdatasystems.com",
    "www.tariffdatasystems.com", "glenrate.com", "www.glenrate.com",
    "paramounttariff.com", "www.paramounttariff.com", "firstbaytariff.com",
    "www.firstbaytariff.com", "login2.etariff.com",
}

_thread_local = threading.local()


@dataclass(frozen=True)
class EntitySource:
    entity_class: str
    organization_no: str
    legal_name: str
    trade_name: str
    active: bool
    tariff_url: str
    directory_url: str
    directory_sha256: str


@dataclass(frozen=True)
class ExtractedTerm:
    rule_type: str
    term_kind: str
    amount_value: str | None
    currency: str | None
    unit: str | None
    quantity_value: str | None
    effective_from: str | None
    effective_to: str | None
    source_version: str | None
    evidence_locator: str
    evidence_excerpt: str
    confidence: float


def utcnow() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def get_session() -> requests.Session:
    session = getattr(_thread_local, "session", None)
    if session is None:
        session = requests.Session()
        session.headers.update({"User-Agent": USER_AGENT, "Accept": "*/*"})
        _thread_local.session = session
    return session


def sha256_bytes(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def canonicalize_url(raw: str, base: str | None = None) -> str | None:
    raw = (raw or "").strip()
    if not raw or raw.startswith(("mailto:", "tel:", "javascript:", "#")):
        return None
    if base:
        raw = urllib.parse.urljoin(base, raw)
    if raw.startswith("//"):
        raw = "https:" + raw
    if not re.match(r"^https?://", raw, re.I):
        raw = "https://" + raw.lstrip("/")
    try:
        p = urllib.parse.urlsplit(raw)
    except ValueError:
        return None
    if not p.hostname:
        return None
    scheme = p.scheme.lower()
    host = p.hostname.lower()
    port = f":{p.port}" if p.port and not (
        (scheme == "https" and p.port == 443) or (scheme == "http" and p.port == 80)
    ) else ""
    path = re.sub(r"/{2,}", "/", p.path or "/")
    query = urllib.parse.parse_qsl(p.query, keep_blank_values=True)
    query = [(k, v) for k, v in query if k.lower() not in {
        "utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content",
        "fbclid", "gclid",
    }]
    query_s = urllib.parse.urlencode(query, doseq=True)
    return urllib.parse.urlunsplit((scheme, host + port, path, query_s, ""))


def same_site(a: str, b: str) -> bool:
    ha = (urllib.parse.urlsplit(a).hostname or "").lower().removeprefix("www.")
    hb = (urllib.parse.urlsplit(b).hostname or "").lower().removeprefix("www.")
    return ha == hb or ha.endswith("." + hb) or hb.endswith("." + ha)


def likely_tariff_link(url: str, anchor_text: str = "", depth: int = 0) -> bool:
    path = urllib.parse.urlsplit(url).path.lower()
    ext = Path(path).suffix.lower()
    if ext in SKIP_EXT:
        return False
    if ext in SUPPORTED_EXT:
        return True
    hay = f"{url} {anchor_text}"
    if TARIFFISH.search(hay):
        return True
    # At the tariff-location landing page, permit one level of same-site navigation.
    return depth == 0 and ext in {"", ".aspx", ".asp", ".php"}


def parse_date(value: str) -> str | None:
    value = re.sub(r"\s+", " ", value.strip()).replace(",", "")
    for fmt in (
        "%Y-%m-%d", "%Y%m%d", "%m/%d/%Y", "%m/%d/%y",
        "%B %d %Y", "%b %d %Y", "%d %B %Y", "%d %b %Y",
        "%d%b%Y", "%d-%b-%Y",
    ):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def detect_effective_dates(text: str) -> tuple[str | None, str | None]:
    head = text[:15000]
    start = end = None
    for pat in DATE_PATTERNS:
        m = pat.search(head)
        if m:
            start = parse_date(m.group("date"))
            if start:
                break
    for pat in END_DATE_PATTERNS:
        m = pat.search(head)
        if m:
            end = parse_date(m.group("date"))
            if end:
                break
    return start, end


def detect_source_version(text: str, url: str) -> str | None:
    probe = (url + "\n" + text[:10000])

    tariff = re.search(
        r"tariff\s+(?:no\.?|number)\s*[:#]?\s*([A-Za-z0-9.\-_]+)",
        probe,
        re.I,
    )
    qualifiers = []
    for label, pattern in (
        ("AMD", r"amendment\s+(?:no\.?|number)?\s*[:#]?\s*([A-Za-z0-9.\-_]+)"),
        ("REV", r"(?:revision|rev\.)\s*[:#]?\s*([A-Za-z0-9.\-_]+)"),
        ("SUP", r"supplement\s+(?:no\.?|number)?\s*[:#]?\s*([A-Za-z0-9.\-_]+)"),
    ):
        match = re.search(pattern, probe, re.I)
        if match:
            qualifier = match.group(1).strip().rstrip(".,;:")
            qualifiers.append(f"{label}:{qualifier}")

    if tariff:
        tariff_id = tariff.group(1).strip().rstrip(".,;:")
        parts = [tariff_id, *qualifiers]
        return "|".join(parts)[:120]
    if qualifiers:
        return "|".join(qualifiers)[:120]

    # Useful fallback for filenames such as Tariff_Rev-6_effective-04_01_2022.pdf.
    name = Path(urllib.parse.urlsplit(url).path).name
    if name and any(token in name.lower() for token in ("tariff", "rate", "schedule")):
        return name[:120]
    return None


def fetch_bytes(url: str, timeout: int, max_bytes: int) -> tuple[bytes, requests.Response]:
    session = get_session()
    with session.get(url, timeout=timeout, allow_redirects=True, stream=True) as resp:
        resp.raise_for_status()
        chunks: list[bytes] = []
        total = 0
        for chunk in resp.iter_content(chunk_size=1024 * 128):
            if not chunk:
                continue
            total += len(chunk)
            if total > max_bytes:
                raise ValueError(f"response exceeds max_bytes={max_bytes}")
            chunks.append(chunk)
        return b"".join(chunks), resp


def text_from_pdf(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    parts: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception as exc:  # corrupt single page should not lose whole snapshot
            txt = f"[PAGE {idx} extraction error: {type(exc).__name__}]"
        parts.append(f"\n[PAGE {idx}]\n{txt}")
    return "\n".join(parts)


def text_from_xlsx(raw: bytes) -> str:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"[SHEET {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v not in (None, "")]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


def extract_text_and_links(
    raw: bytes,
    content_type: str,
    url: str,
) -> tuple[str, list[tuple[str, str]], str, str | None]:
    ctype = (content_type or "").lower()
    path = urllib.parse.urlsplit(url).path.lower()
    ext = Path(path).suffix.lower()
    links: list[tuple[str, str]] = []
    title = None

    if "pdf" in ctype or ext == ".pdf":
        return text_from_pdf(raw), links, "pdf", title

    if ext in {".xlsx", ".xlsm"} or "spreadsheet" in ctype:
        return text_from_xlsx(raw), links, "xlsx", title

    decoded = raw.decode("utf-8", errors="replace")

    if "html" in ctype or ext in {".html", ".htm", ".asp", ".aspx", ".php", ""}:
        soup = BeautifulSoup(decoded, "html.parser")
        if soup.title:
            title = soup.title.get_text(" ", strip=True)[:300]
        for a in soup.find_all("a", href=True):
            href = canonicalize_url(a.get("href"), url)
            if href:
                links.append((href, a.get_text(" ", strip=True)[:300]))
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return soup.get_text("\n", strip=True), links, "html", title

    return decoded, links, "text", title


def split_evidence_blocks(text: str) -> Iterator[tuple[str, str]]:
    """Yield locator + compact paragraph/line blocks with stable-ish evidence anchors."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    buf: list[str] = []
    start = 1
    for idx, line in enumerate(lines, start=1):
        if not line:
            if buf:
                yield f"lines:{start}-{idx-1}", " ".join(buf)[:4000]
                buf = []
            start = idx + 1
            continue
        if not buf:
            start = idx
        buf.append(line)
        if len(" ".join(buf)) > 1800:
            yield f"lines:{start}-{idx}", " ".join(buf)[:4000]
            buf = []
            start = idx + 1
    if buf:
        yield f"lines:{start}-{len(lines)}", " ".join(buf)[:4000]


def classify_rules(block: str) -> list[str]:
    return [name for name, pat in RULE_PATTERNS if pat.search(block)]


def extract_terms(text: str, url: str) -> list[ExtractedTerm]:
    doc_start, doc_end = detect_effective_dates(text)
    source_version = detect_source_version(text, url)
    terms: list[ExtractedTerm] = []
    seen: set[tuple] = set()

    for locator, block in split_evidence_blocks(text):
        rules = classify_rules(block)
        if not rules:
            continue

        local_start, local_end = detect_effective_dates(block)
        effective_from = local_start or doc_start
        effective_to = local_end or doc_end

        money_matches = list(MONEY_RE.finditer(block))
        pct_matches = list(PERCENT_RE.finditer(block))
        free_match = FREE_DAYS_RE.search(block)

        for rule in rules:
            for m in money_matches:
                unit = (m.group("unit") or "").strip() or None
                key = (rule, "money", m.group("amount"), "USD", unit, locator, block[:800])
                if key in seen:
                    continue
                seen.add(key)
                terms.append(ExtractedTerm(
                    rule_type=rule,
                    term_kind="money",
                    amount_value=m.group("amount").replace(",", ""),
                    currency="USD",
                    unit=unit,
                    quantity_value=None,
                    effective_from=effective_from,
                    effective_to=effective_to,
                    source_version=source_version,
                    evidence_locator=locator,
                    evidence_excerpt=block[:2000],
                    confidence=0.94 if TARIFFISH.search(block) else 0.80,
                ))

            for m in pct_matches:
                key = (rule, "percent", m.group("amount"), locator, block[:800])
                if key in seen:
                    continue
                seen.add(key)
                terms.append(ExtractedTerm(
                    rule_type=rule,
                    term_kind="percent",
                    amount_value=m.group("amount"),
                    currency=None,
                    unit=(m.group("unit") or "%").strip(),
                    quantity_value=None,
                    effective_from=effective_from,
                    effective_to=effective_to,
                    source_version=source_version,
                    evidence_locator=locator,
                    evidence_excerpt=block[:2000],
                    confidence=0.90,
                ))

            if rule == "free_time" and free_match:
                qty = free_match.group("qty")
                key = (rule, "quantity", qty, locator, block[:800])
                if key not in seen:
                    seen.add(key)
                    terms.append(ExtractedTerm(
                        rule_type=rule,
                        term_kind="quantity",
                        amount_value=None,
                        currency=None,
                        unit="days",
                        quantity_value=qty,
                        effective_from=effective_from,
                        effective_to=effective_to,
                        source_version=source_version,
                        evidence_locator=locator,
                        evidence_excerpt=block[:2000],
                        confidence=0.96,
                    ))

            # Preserve non-monetary rule text when it has rule semantics but no amount.
            if not money_matches and not pct_matches and not (rule == "free_time" and free_match):
                key = (rule, "rule_text", locator, block[:800])
                if key not in seen:
                    seen.add(key)
                    terms.append(ExtractedTerm(
                        rule_type=rule,
                        term_kind="rule_text",
                        amount_value=None,
                        currency=None,
                        unit=None,
                        quantity_value=None,
                        effective_from=effective_from,
                        effective_to=effective_to,
                        source_version=source_version,
                        evidence_locator=locator,
                        evidence_excerpt=block[:2000],
                        confidence=0.78,
                    ))

    return terms


def init_db(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    schema = Path(__file__).with_name("schema.sql").read_text(encoding="utf-8")
    conn.executescript(schema)
    conn.commit()
    return conn


def parse_directory_report(
    entity_class: str,
    active: bool,
    url: str,
    timeout: int,
    max_bytes: int,
) -> tuple[list[EntitySource], str]:
    raw, resp = fetch_bytes(url, timeout, max_bytes)
    digest = sha256_bytes(raw)
    soup = BeautifulSoup(raw, "html.parser")
    rows: list[EntitySource] = []

    for tr in soup.find_all("tr"):
        tds = tr.find_all("td")
        if len(tds) < 3:
            continue
        cells = [td.get_text(" ", strip=True) for td in tds]
        org = cells[0].strip()
        if not re.fullmatch(r"\d{4,8}", org):
            continue
        legal = cells[1].strip() if len(cells) > 1 else ""
        trade = cells[2].strip() if len(cells) > 3 else ""
        loc_cell = tds[-1]

        locs: list[str] = []
        for a in loc_cell.find_all("a", href=True):
            href = canonicalize_url(a["href"], url)
            if href:
                locs.append(href)

        if not locs:
            loc_text = loc_cell.get_text(" ", strip=True)
            candidates = re.findall(
                r"(?:https?://)?(?:www\.)?[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
                r"(?:/[A-Za-z0-9_?&=%.+~:/#@\-]*)?",
                loc_text,
            )
            for candidate in candidates:
                href = canonicalize_url(candidate)
                if href:
                    locs.append(href)

        for tariff_url in sorted(set(locs)):
            rows.append(EntitySource(
                entity_class=entity_class,
                organization_no=org,
                legal_name=legal or "(unknown)",
                trade_name=trade,
                active=active,
                tariff_url=tariff_url,
                directory_url=url,
                directory_sha256=digest,
            ))

    return rows, digest


def enumerate_fmc(timeout: int = 30, max_bytes: int = 15_000_000) -> list[EntitySource]:
    all_rows: list[EntitySource] = []
    for entity_class, active, url in DIRECTORY_REPORTS:
        try:
            rows, _ = parse_directory_report(entity_class, active, url, timeout, max_bytes)
            all_rows.extend(rows)
        except Exception as exc:
            print(f"WARN directory fetch failed {url}: {exc}", file=sys.stderr)

    # Deduplicate identical current/history references while preserving active/inactive status.
    dedup: dict[tuple, EntitySource] = {}
    for row in all_rows:
        key = (
            row.entity_class, row.organization_no, row.legal_name,
            row.trade_name, row.active, row.tariff_url,
        )
        dedup[key] = row
    return sorted(
        dedup.values(),
        key=lambda r: (r.entity_class, r.organization_no, r.tariff_url, not r.active),
    )


def persist_source(conn: sqlite3.Connection, src: EntitySource) -> int:
    now = utcnow()
    conn.execute(
        """INSERT OR IGNORE INTO entities
           (entity_class, organization_no, legal_name, trade_name, active)
           VALUES (?, ?, ?, ?, ?)""",
        (src.entity_class, src.organization_no, src.legal_name, src.trade_name, int(src.active)),
    )
    row = conn.execute(
        """SELECT id FROM entities
           WHERE entity_class=? AND organization_no=? AND legal_name=?
             AND trade_name=? AND active=?""",
        (src.entity_class, src.organization_no, src.legal_name, src.trade_name, int(src.active)),
    ).fetchone()
    assert row
    entity_id = int(row[0])

    conn.execute(
        """INSERT OR IGNORE INTO tariff_locations
           (entity_id, canonical_url, directory_url, directory_snapshot_sha256,
            first_seen_at, last_seen_at)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (entity_id, src.tariff_url, src.directory_url, src.directory_sha256, now, now),
    )
    conn.execute(
        """UPDATE tariff_locations
           SET last_seen_at=?, directory_url=?, directory_snapshot_sha256=?
           WHERE entity_id=? AND canonical_url=?""",
        (now, src.directory_url, src.directory_sha256, entity_id, src.tariff_url),
    )
    loc = conn.execute(
        "SELECT id FROM tariff_locations WHERE entity_id=? AND canonical_url=?",
        (entity_id, src.tariff_url),
    ).fetchone()
    assert loc
    conn.commit()
    return int(loc[0])


def write_blob(blob_root: Path, digest: str, raw: bytes) -> str:
    rel = Path("blobs") / "sha256" / digest[:2] / digest
    path = blob_root.parent / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        tmp = path.with_suffix(".tmp")
        tmp.write_bytes(raw)
        os.replace(tmp, path)
    return str(rel)


def portal_note(url: str) -> str | None:
    host = (urllib.parse.urlsplit(url).hostname or "").lower()
    if host in PORTAL_DOMAINS:
        return host
    return None


def crawl_one_location(
    src: EntitySource,
    out_dir: Path,
    max_depth: int,
    timeout: int,
    max_bytes: int,
    max_pages: int,
) -> dict:
    """Network crawl one tariff location into an isolated shard DB write payload."""
    queue: list[tuple[str, int]] = [(src.tariff_url, 0)]
    seen: set[str] = set()
    snapshots: list[dict] = []
    errors: list[dict] = []
    terms_payload: list[tuple[int, ExtractedTerm]] = []

    while queue and len(snapshots) < max_pages:
        requested_url, depth = queue.pop(0)
        requested_url = canonicalize_url(requested_url) or requested_url
        if requested_url in seen:
            continue
        seen.add(requested_url)

        try:
            raw, resp = fetch_bytes(requested_url, timeout, max_bytes)
            final_url = canonicalize_url(resp.url) or resp.url
            digest = sha256_bytes(raw)
            ctype = resp.headers.get("content-type", "").split(";")[0].strip().lower()
            text = ""
            links: list[tuple[str, str]] = []
            parse_kind = "binary"
            title = None
            parser_status = "snapshotted"
            try:
                text, links, parse_kind, title = extract_text_and_links(raw, ctype, final_url)
                parser_status = f"parsed:{parse_kind}"
            except Exception as exc:
                parser_status = f"snapshot_only:{type(exc).__name__}"
                errors.append({
                    "url": final_url,
                    "stage": "parse",
                    "error_type": type(exc).__name__,
                    "detail": str(exc)[:2000],
                })

            doc_start, doc_end = detect_effective_dates(text) if text else (None, None)
            version = detect_source_version(text, final_url) if text else None
            blob_relpath = write_blob(out_dir / "blobs", digest, raw)

            snap_idx = len(snapshots)
            snapshots.append({
                "requested_url": requested_url,
                "final_url": final_url,
                "fetched_at": utcnow(),
                "http_status": resp.status_code,
                "content_type": ctype,
                "byte_count": len(raw),
                "sha256": digest,
                "blob_relpath": blob_relpath,
                "parser_status": parser_status,
                "title": title,
                "source_version": version,
                "effective_from": doc_start,
                "effective_to": doc_end,
            })

            if text:
                for term in extract_terms(text, final_url):
                    terms_payload.append((snap_idx, term))

            if depth < max_depth and links:
                for href, anchor in links[:400]:
                    if href in seen:
                        continue
                    if not same_site(src.tariff_url, href):
                        continue
                    if likely_tariff_link(href, anchor, depth):
                        queue.append((href, depth + 1))

            # Known tariff publishers frequently require client-side search/login.
            if depth == 0 and portal_note(final_url) and len(links) == 0:
                errors.append({
                    "url": final_url,
                    "stage": "portal_adapter",
                    "error_type": "PortalAdapterNeeded",
                    "detail": (
                        f"Known publisher portal {portal_note(final_url)} exposed no crawlable "
                        "tariff links in static HTML; preserve landing snapshot and route to "
                        "publisher-specific browser adapter."
                    ),
                })

        except Exception as exc:
            errors.append({
                "url": requested_url,
                "stage": "fetch",
                "error_type": type(exc).__name__,
                "detail": str(exc)[:2000],
            })

    return {
        "source": src,
        "snapshots": snapshots,
        "terms": terms_payload,
        "errors": errors,
        "seen_count": len(seen),
    }


def persist_crawl_result(conn: sqlite3.Connection, result: dict, out_dir: Path) -> dict:
    src: EntitySource = result["source"]
    loc_id = persist_source(conn, src)
    snap_ids: list[int] = []
    now = utcnow()

    for snap in result["snapshots"]:
        conn.execute(
            """INSERT OR IGNORE INTO snapshots
               (tariff_location_id, requested_url, final_url, fetched_at, http_status,
                content_type, byte_count, sha256, blob_relpath, parser_status, title,
                source_version, effective_from, effective_to)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                loc_id, snap["requested_url"], snap["final_url"], snap["fetched_at"],
                snap["http_status"], snap["content_type"], snap["byte_count"],
                snap["sha256"], snap["blob_relpath"], snap["parser_status"],
                snap["title"], snap["source_version"], snap["effective_from"],
                snap["effective_to"],
            ),
        )
        row = conn.execute(
            """SELECT id FROM snapshots
               WHERE tariff_location_id=? AND requested_url=? AND sha256=?""",
            (loc_id, snap["requested_url"], snap["sha256"]),
        ).fetchone()
        assert row
        snap_ids.append(int(row[0]))

    for snap_idx, term in result["terms"]:
        if snap_idx >= len(snap_ids):
            continue
        conn.execute(
            """INSERT INTO terms
               (snapshot_id, entity_class, organization_no, legal_name, rule_type,
                term_kind, amount_value, currency, unit, quantity_value,
                effective_from, effective_to, source_version, evidence_locator,
                evidence_excerpt, confidence, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                snap_ids[snap_idx], src.entity_class, src.organization_no,
                src.legal_name, term.rule_type, term.term_kind, term.amount_value,
                term.currency, term.unit, term.quantity_value, term.effective_from,
                term.effective_to, term.source_version, term.evidence_locator,
                term.evidence_excerpt, term.confidence, now,
            ),
        )

    for err in result["errors"]:
        conn.execute(
            """INSERT INTO crawl_errors
               (tariff_location_id, url, occurred_at, stage, error_type, detail)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (loc_id, err["url"], now, err["stage"], err["error_type"], err["detail"]),
        )

    conn.commit()
    return {
        "snapshots": len(result["snapshots"]),
        "terms": len(result["terms"]),
        "errors": len(result["errors"]),
    }


def shard_sources(
    sources: list[EntitySource],
    shard_index: int,
    shard_count: int,
) -> list[EntitySource]:
    if shard_count <= 1:
        return sources
    out = []
    for src in sources:
        stable = hashlib.sha256(
            f"{src.entity_class}|{src.organization_no}|{src.tariff_url}".encode()
        ).digest()
        bucket = int.from_bytes(stable[:8], "big") % shard_count
        if bucket == shard_index:
            out.append(src)
    return out


def command_enumerate(args: argparse.Namespace) -> int:
    rows = enumerate_fmc(args.timeout, args.max_bytes)
    payload = [
        {
            "entity_class": r.entity_class,
            "organization_no": r.organization_no,
            "legal_name": r.legal_name,
            "trade_name": r.trade_name,
            "active": r.active,
            "tariff_url": r.tariff_url,
            "directory_url": r.directory_url,
            "directory_sha256": r.directory_sha256,
        }
        for r in rows
    ]
    if args.format == "json":
        print(json.dumps(payload, indent=2))
    else:
        writer = csv.DictWriter(sys.stdout, fieldnames=list(payload[0]) if payload else [])
        if payload:
            writer.writeheader()
            writer.writerows(payload)
    return 0


def command_crawl(args: argparse.Namespace) -> int:
    out_dir = Path(args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    db_path = Path(args.db)
    if not db_path.is_absolute():
        db_path = out_dir / db_path
    conn = init_db(db_path)

    started = utcnow()
    cur = conn.execute(
        """INSERT INTO crawl_runs
           (started_at, shard_index, shard_count, status)
           VALUES (?, ?, ?, ?)""",
        (started, args.shard_index, args.shard_count, "running"),
    )
    run_id = int(cur.lastrowid)
    conn.commit()

    sources = enumerate_fmc(args.timeout, args.max_bytes)
    sources = shard_sources(sources, args.shard_index, args.shard_count)
    if args.max_locations:
        sources = sources[: args.max_locations]

    stats = {
        "directory_locations": len(sources),
        "snapshots": 0,
        "terms": 0,
        "errors": 0,
        "known_portal_locations": sum(bool(portal_note(s.tariff_url)) for s in sources),
    }

    # Persist directory rows before crawling so inaccessible locations still exist in ledger.
    for src in sources:
        persist_source(conn, src)

    print(
        f"FMC ledger crawl: {len(sources)} tariff locations "
        f"(shard {args.shard_index}/{args.shard_count})",
        file=sys.stderr,
    )

    with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = [
            pool.submit(
                crawl_one_location,
                src,
                out_dir,
                args.max_depth,
                args.timeout,
                args.max_bytes,
                args.max_pages_per_location,
            )
            for src in sources
        ]
        for idx, fut in enumerate(concurrent.futures.as_completed(futures), start=1):
            try:
                result = fut.result()
                delta = persist_crawl_result(conn, result, out_dir)
                for k in ("snapshots", "terms", "errors"):
                    stats[k] += delta[k]
            except Exception as exc:
                stats["errors"] += 1
                print(f"ERROR worker result: {exc}", file=sys.stderr)

            if idx % 25 == 0 or idx == len(futures):
                print(
                    f"progress {idx}/{len(futures)} "
                    f"snapshots={stats['snapshots']} terms={stats['terms']} "
                    f"errors={stats['errors']}",
                    file=sys.stderr,
                )

    stats["finished_at"] = utcnow()
    conn.execute(
        """UPDATE crawl_runs
           SET finished_at=?, status=?, stats_json=?
           WHERE id=?""",
        (stats["finished_at"], "complete", json.dumps(stats, sort_keys=True), run_id),
    )
    conn.commit()

    stats_path = out_dir / "stats.json"
    stats_path.write_text(json.dumps(stats, indent=2, sort_keys=True) + "\n")
    print(json.dumps(stats, indent=2, sort_keys=True))
    return 0


def copy_rows(dst: sqlite3.Connection, src: sqlite3.Connection) -> None:
    """Merge a shard DB by natural keys while remapping IDs."""
    entity_map: dict[int, int] = {}
    loc_map: dict[int, int] = {}
    snap_map: dict[int, int] = {}

    for row in src.execute(
        "SELECT id, entity_class, organization_no, legal_name, trade_name, active FROM entities"
    ):
        old_id, cls, org, legal, trade, active = row
        dst.execute(
            """INSERT OR IGNORE INTO entities
               (entity_class, organization_no, legal_name, trade_name, active)
               VALUES (?, ?, ?, ?, ?)""",
            (cls, org, legal, trade, active),
        )
        new = dst.execute(
            """SELECT id FROM entities WHERE entity_class=? AND organization_no=?
               AND legal_name=? AND trade_name=? AND active=?""",
            (cls, org, legal, trade, active),
        ).fetchone()
        assert new
        entity_map[int(old_id)] = int(new[0])

    for row in src.execute(
        """SELECT id, entity_id, canonical_url, directory_url,
                  directory_snapshot_sha256, first_seen_at, last_seen_at
           FROM tariff_locations"""
    ):
        old_id, old_entity, url, durl, dsha, first_seen, last_seen = row
        new_entity = entity_map[int(old_entity)]
        dst.execute(
            """INSERT OR IGNORE INTO tariff_locations
               (entity_id, canonical_url, directory_url, directory_snapshot_sha256,
                first_seen_at, last_seen_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (new_entity, url, durl, dsha, first_seen, last_seen),
        )
        dst.execute(
            """UPDATE tariff_locations
               SET last_seen_at=MAX(last_seen_at, ?),
                   directory_url=?, directory_snapshot_sha256=?
               WHERE entity_id=? AND canonical_url=?""",
            (last_seen, durl, dsha, new_entity, url),
        )
        new = dst.execute(
            "SELECT id FROM tariff_locations WHERE entity_id=? AND canonical_url=?",
            (new_entity, url),
        ).fetchone()
        assert new
        loc_map[int(old_id)] = int(new[0])

    for row in src.execute(
        """SELECT id, tariff_location_id, requested_url, final_url, fetched_at,
                  http_status, content_type, byte_count, sha256, blob_relpath,
                  parser_status, title, source_version, effective_from, effective_to
           FROM snapshots"""
    ):
        (
            old_id, old_loc, requested, final, fetched, status, ctype, byte_count,
            digest, blob, parser_status, title, version, eff_from, eff_to,
        ) = row
        new_loc = loc_map[int(old_loc)]
        dst.execute(
            """INSERT OR IGNORE INTO snapshots
               (tariff_location_id, requested_url, final_url, fetched_at, http_status,
                content_type, byte_count, sha256, blob_relpath, parser_status, title,
                source_version, effective_from, effective_to)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                new_loc, requested, final, fetched, status, ctype, byte_count,
                digest, blob, parser_status, title, version, eff_from, eff_to,
            ),
        )
        new = dst.execute(
            """SELECT id FROM snapshots WHERE tariff_location_id=? AND requested_url=?
               AND sha256=?""",
            (new_loc, requested, digest),
        ).fetchone()
        assert new
        snap_map[int(old_id)] = int(new[0])

    for row in src.execute(
        """SELECT snapshot_id, entity_class, organization_no, legal_name, rule_type,
                  term_kind, amount_value, currency, unit, quantity_value,
                  effective_from, effective_to, source_version, evidence_locator,
                  evidence_excerpt, confidence, created_at
           FROM terms"""
    ):
        old_snap = int(row[0])
        if old_snap not in snap_map:
            continue
        dst.execute(
            """INSERT INTO terms
               (snapshot_id, entity_class, organization_no, legal_name, rule_type,
                term_kind, amount_value, currency, unit, quantity_value,
                effective_from, effective_to, source_version, evidence_locator,
                evidence_excerpt, confidence, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (snap_map[old_snap], *row[1:]),
        )

    for row in src.execute(
        "SELECT tariff_location_id, url, occurred_at, stage, error_type, detail FROM crawl_errors"
    ):
        old_loc = row[0]
        new_loc = loc_map.get(int(old_loc)) if old_loc is not None else None
        dst.execute(
            """INSERT INTO crawl_errors
               (tariff_location_id, url, occurred_at, stage, error_type, detail)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (new_loc, *row[1:]),
        )
    dst.commit()


def command_merge(args: argparse.Namespace) -> int:
    root = Path(args.input_root)
    candidates = sorted(root.rglob("ledger.sqlite"))
    if not candidates:
        raise SystemExit(f"no ledger.sqlite files found under {root}")
    db_path = Path(args.db)
    conn = init_db(db_path)
    for candidate in candidates:
        print(f"merging {candidate}", file=sys.stderr)
        src = sqlite3.connect(candidate)
        copy_rows(conn, src)
        src.close()
    command_summary(argparse.Namespace(db=str(db_path)))
    return 0


def command_summary(args: argparse.Namespace) -> int:
    conn = sqlite3.connect(args.db)
    queries = {
        "entities": "SELECT COUNT(*) FROM entities",
        "tariff_locations": "SELECT COUNT(*) FROM tariff_locations",
        "snapshots": "SELECT COUNT(*) FROM snapshots",
        "unique_snapshot_hashes": "SELECT COUNT(DISTINCT sha256) FROM snapshots WHERE sha256 IS NOT NULL",
        "terms": "SELECT COUNT(*) FROM terms",
        "errors": "SELECT COUNT(*) FROM crawl_errors",
        "rules": "SELECT COUNT(DISTINCT rule_type) FROM terms",
        "term_effective_dated_terms": "SELECT COUNT(*) FROM terms WHERE effective_from IS NOT NULL",
        "effective_dated_terms": """SELECT COUNT(*) FROM terms t JOIN snapshots s ON s.id=t.snapshot_id
                                      WHERE COALESCE(t.effective_from, s.effective_from) IS NOT NULL""",
        "source_versioned_terms": """SELECT COUNT(*) FROM terms t JOIN snapshots s ON s.id=t.snapshot_id
                                      WHERE COALESCE(t.source_version, s.source_version) IS NOT NULL""",
    }
    stats = {name: conn.execute(sql).fetchone()[0] for name, sql in queries.items()}
    stats["effective_date_coverage_pct"] = round(
        (100.0 * stats["effective_dated_terms"] / stats["terms"]) if stats["terms"] else 0.0,
        2,
    )
    stats["source_version_coverage_pct"] = round(
        (100.0 * stats["source_versioned_terms"] / stats["terms"]) if stats["terms"] else 0.0,
        2,
    )
    stats["entity_classes"] = conn.execute(
        """SELECT entity_class, COUNT(*) c FROM entities
           GROUP BY entity_class ORDER BY c DESC"""
    ).fetchall()
    stats["top_rule_types"] = conn.execute(
        """SELECT rule_type, COUNT(*) c FROM terms
           GROUP BY rule_type ORDER BY c DESC LIMIT 30"""
    ).fetchall()
    stats["currencies"] = conn.execute(
        """SELECT COALESCE(currency, '(none)') currency, COUNT(*) c FROM terms
           GROUP BY 1 ORDER BY c DESC LIMIT 30"""
    ).fetchall()
    stats["parser_statuses"] = conn.execute(
        """SELECT parser_status, COUNT(*) c FROM snapshots
           GROUP BY parser_status ORDER BY c DESC LIMIT 40"""
    ).fetchall()
    stats["top_publishers"] = conn.execute(
        """SELECT lower(
                 replace(
                   replace(substr(canonical_url, instr(canonical_url, '//')+2), 'www.', ''),
                   '/', ''
                 )
               ) host,
               COUNT(*) c
           FROM tariff_locations
           GROUP BY host ORDER BY c DESC LIMIT 40"""
    ).fetchall()
    stats["error_types"] = conn.execute(
        """SELECT stage || ':' || error_type, COUNT(*) c FROM crawl_errors
           GROUP BY 1 ORDER BY c DESC LIMIT 40"""
    ).fetchall()
    stats["error_domains"] = conn.execute(
        """SELECT lower(
                 replace(
                   replace(substr(url, instr(url, '//')+2), 'www.', ''),
                   '/', ''
                 )
               ) host,
               stage || ':' || error_type reason,
               COUNT(*) c
           FROM crawl_errors
           WHERE instr(url, '//') > 0
           GROUP BY host, reason
           ORDER BY c DESC LIMIT 80"""
    ).fetchall()
    stats["unresolved_publishers"] = conn.execute(
        """SELECT lower(
                 replace(
                   replace(substr(url, instr(url, '//')+2), 'www.', ''),
                   '/', ''
                 )
               ) host,
               COUNT(*) c
           FROM crawl_errors
           WHERE stage='publisher_adapter'
             AND error_type='EntityTariffNotResolved'
             AND instr(url, '//') > 0
           GROUP BY host ORDER BY c DESC LIMIT 50"""
    ).fetchall()
    print(json.dumps(stats, indent=2))
    return 0



def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="FMC tariff rule ledger")
    sub = p.add_subparsers(dest="command", required=True)

    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--timeout", type=int, default=30)
    common.add_argument("--max-bytes", type=int, default=50_000_000)

    e = sub.add_parser("enumerate", parents=[common])
    e.add_argument("--format", choices=["json", "csv"], default="json")
    e.set_defaults(func=command_enumerate)

    c = sub.add_parser("crawl", parents=[common])
    c.add_argument("--out", default=".fmc_tariff_ledger")
    c.add_argument("--db", default="ledger.sqlite")
    c.add_argument("--workers", type=int, default=8)
    c.add_argument("--max-depth", type=int, default=2)
    c.add_argument("--max-pages-per-location", type=int, default=60)
    c.add_argument("--max-locations", type=int, default=0)
    c.add_argument("--shard-index", type=int, default=0)
    c.add_argument("--shard-count", type=int, default=1)
    c.set_defaults(func=command_crawl)

    m = sub.add_parser("merge")
    m.add_argument("--input-root", required=True)
    m.add_argument("--db", required=True)
    m.set_defaults(func=command_merge)

    s = sub.add_parser("summary")
    s.add_argument("--db", required=True)
    s.set_defaults(func=command_summary)
    return p


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    if getattr(args, "shard_index", 0) < 0:
        parser.error("--shard-index must be >= 0")
    if getattr(args, "shard_count", 1) < 1:
        parser.error("--shard-count must be >= 1")
    if getattr(args, "shard_index", 0) >= getattr(args, "shard_count", 1):
        parser.error("--shard-index must be < --shard-count")
    return int(args.func(args))


if __name__ == "__main__":
    raise SystemExit(main())
